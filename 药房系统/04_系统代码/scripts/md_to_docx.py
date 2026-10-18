# -*- coding: utf-8 -*-
"""
轻量 Markdown → Word（.docx）转换器（够用即可：标题 / 表格 / 列表 / 代码块 / 引用 / 粗体）。

用法：python scripts/md_to_docx.py <输入.md> <输出.docx>
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

CN_FONT = "微软雅黑"
CN_MONO = "Consolas"


def set_font(run, name=CN_FONT, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_inline(p, text, base_size=10.5, base_bold=False):
    """处理 **粗体** 与 `代码`（用 split 分隔符交替扫描）"""
    state_b = False
    state_c = False
    for tok in re.split(r"(\*\*|`)", text):
        if tok == "**":
            state_b = not state_b
            continue
        if tok == "`":
            state_c = not state_c
            continue
        if not tok:
            continue
        r = p.add_run(tok)
        set_font(r, CN_MONO if state_c else CN_FONT, base_size, base_bold or state_b,
                 (0xC0, 0x39, 0x2B) if state_c else None)


def convert(md_path, docx_path):
    lines = open(md_path, encoding="utf-8").read().splitlines()
    doc = Document()
    # 正文样式
    st = doc.styles["Normal"]
    st.font.name = CN_FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()

        # 代码块
        if s.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(buf))
            set_font(r, CN_MONO, 9.5)
            p.paragraph_format.left_indent = Pt(18)
            continue

        # 空行
        if not s:
            i += 1
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}", s):
            doc.add_paragraph("─" * 40)
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            lvl = min(len(m.group(1)), 4)
            p = doc.add_heading(level=lvl)
            r = p.add_run(m.group(2))
            set_font(r, CN_FONT, [16, 14, 12, 11][lvl - 1], True)
            i += 1
            continue

        # 表格
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            header = [c.strip() for c in rows[0].strip("|").split("|")]
            body = [[c.strip() for c in r.strip("|").split("|")] for r in rows[2:]]
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Light Grid Accent 1"
            for j, h in enumerate(header):
                cell = t.rows[0].cells[j]
                cell.text = ""
                add_inline(cell.paragraphs[0], h, 10, True)
            for row in body:
                cells = t.add_row().cells
                for j, v in enumerate(row[:len(header)]):
                    cells[j].text = ""
                    add_inline(cells[j].paragraphs[0], v.replace("<br>", " / "), 9.5)
            doc.add_paragraph()
            continue

        # 引用
        if s.startswith(">"):
            p = doc.add_paragraph()
            add_inline(p, s.lstrip("> ").strip(), 10)
            p.paragraph_format.left_indent = Pt(18)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x57, 0x60, 0x6A)
            i += 1
            continue

        # 列表
        m = re.match(r"^[-*]\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, m.group(1))
            i += 1
            continue
        m = re.match(r"^\d+\.\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, m.group(1))
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_inline(p, s)
        i += 1

    doc.save(docx_path)
    return docx_path


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        raise SystemExit(1)
    out = convert(sys.argv[1], sys.argv[2])
    print("✔ 已生成：", out)
