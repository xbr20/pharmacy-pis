# -*- coding: utf-8 -*-
"""数据集 Excel 自检：Sheet 清单、说明页内容、表头样例、空表检查"""
import os
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "data", "药房信息系统_数据集.xlsx")
OUT = os.path.join(ROOT, "data", "_自检报告.txt")

lines = []
wb = load_workbook(XLSX)
lines.append("Sheet 数：%d" % len(wb.sheetnames))
lines.append("顺序：" + " | ".join(wb.sheetnames))

ws = wb["00_数据集说明"]
lines.append("\n[00_数据集说明]")
for row in ws.iter_rows(min_row=1, max_row=22, values_only=True):
    if any(v is not None for v in row):
        lines.append("  " + " | ".join("" if v is None else str(v) for v in row))

ws = wb["01_数据字典"]
lines.append("\n[01_数据字典] 行数：%d" % ws.max_row)
for row in ws.iter_rows(min_row=1, max_row=5, values_only=True):
    lines.append("  " + " | ".join("" if v is None else str(v) for v in row))

lines.append("\n[各表]")
for name in wb.sheetnames:
    if name.startswith(("00_", "01_")):
        continue
    ws = wb[name]
    lines.append("  %-14s %4d 行 x %2d 列   表头: %s"
                 % (name, ws.max_row - 1, ws.max_column,
                    "、".join(str(c.value) for c in ws[1])[:90]))

open(OUT, "w", encoding="utf-8").write("\n".join(lines))
print("自检报告已写入：", OUT)
print("文件大小：%.1f KB" % (os.path.getsize(XLSX) / 1024))
